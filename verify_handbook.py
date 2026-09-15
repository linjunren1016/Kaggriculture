"""Verify the generated handbook .docx: structure, tables, and text integrity."""
import sys
from docx import Document

path = r"D:\hiwsk\Documents\deepseek工作区\kaggriculture\Kaggriculture队员手册.docx"
doc = Document(path)

heads = [(p.style.name, p.text) for p in doc.paragraphs
         if p.style.name.startswith("Heading") or p.style.name == "Title"]
print(f"paragraphs: {len(doc.paragraphs)}   tables: {len(doc.tables)}")
print("\n--- 标题结构 ---")
for style, text in heads:
    if not text.strip():
        continue
    level = 0 if style == "Title" else int(style.split()[-1])
    print("  " * level + f"[{style}] {text}")

print("\n--- 表格概览 ---")
for i, t in enumerate(doc.tables):
    hdr = " | ".join(c.text.strip() for c in t.rows[0].cells)
    print(f"  T{i}: {len(t.rows)-1} 行数据 | 表头: {hdr[:70]}")

# 关键结论必须存在
must_have = [
    "只看赢 / 输 / 平",
    "全部提交中最新的 2 次",
    "是「全部提交中最后提交的 2 个」",
    "截止前必须有意安排最后两次提交",
    "必须跑 10–20 个随机种子",
    "做规则，不要急着建模",
    "从后往前优化",
]
alltext = "\n".join(p.text for p in doc.paragraphs)
for t in doc.tables:
    for r in t.rows:
        for c in r.cells:
            alltext += "\n" + c.text
print("\n--- 关键结论检查 ---")
missing = []
for k in must_have:
    ok = k in alltext
    print(f"  {'OK ' if ok else 'MISS'}  {k}")
    if not ok:
        missing.append(k)

# 不应再出现对外部资料的依赖
banned = ["The Kaggle Book", "Bojan", "kedou.life", "bilibili", "BruceQD", "待办", "每天要做什么"]
print("\n--- 外部依赖 / 待办 检查（应全部为 0）---")
for b in banned:
    n = alltext.count(b)
    print(f"  {b}: {n}")

print("\nRESULT:", "PASS" if not missing else f"FAIL missing={missing}")
sys.exit(1 if missing else 0)
