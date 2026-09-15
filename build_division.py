"""
生成《人员与 AI 分工》Word 文档。

排版与《Kaggriculture队员手册.docx》保持一致，内容自足。
"""
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

CN_FONT = "微软雅黑"
ACCENT = RGBColor(0x1F, 0x4E, 0x79)
WARN_RED = RGBColor(0xC0, 0x00, 0x00)
HUMAN = RGBColor(0x1F, 0x4E, 0x79)     # 蓝：人做
AI = RGBColor(0x38, 0x76, 0x1D)        # 绿：AI 做


def set_base_style(doc):
    st = doc.styles["Normal"]
    st.font.name = CN_FONT
    st.font.size = Pt(10.5)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), CN_FONT)
    st.paragraph_format.space_after = Pt(4)
    st.paragraph_format.line_spacing = 1.25


def style_heading(doc, name, size, color):
    st = doc.styles[name]
    st.font.name = CN_FONT
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = color
    st.element.rPr.rFonts.set(qn("w:eastAsia"), CN_FONT)


def shade(cell, hex_fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    el = OxmlElement("w:shd")
    el.set(qn("w:val"), "clear")
    el.set(qn("w:fill"), hex_fill)
    tc_pr.append(el)


def h1(doc, t):
    return doc.add_heading(t, level=1)


def h2(doc, t):
    return doc.add_heading(t, level=2)


def p(doc, text, bold=False, color=None, size=None):
    par = doc.add_paragraph()
    run = par.add_run(text)
    run.bold = bold
    if color is not None:
        run.font.color.rgb = color
    if size is not None:
        run.font.size = Pt(size)
    return par


def rich(doc, segments):
    par = doc.add_paragraph()
    for text, bold, color in segments:
        run = par.add_run(text)
        run.bold = bold
        if color is not None:
            run.font.color.rgb = color
    return par


def bullets(doc, items):
    for it in items:
        if isinstance(it, tuple):
            par = doc.add_paragraph(style="List Bullet")
            par.add_run(it[0]).bold = True
            par.add_run(it[1])
        else:
            doc.add_paragraph(it, style="List Bullet")


def warn(doc, text):
    par = doc.add_paragraph()
    run = par.add_run("⚠  " + text)
    run.bold = True
    run.font.color.rgb = WARN_RED
    return par


def table(doc, header, rows, widths=None, header_fill="1F4E79"):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = t.rows[0].cells
    for i, text in enumerate(header):
        hdr[i].text = ""
        run = hdr[i].paragraphs[0].add_run(text)
        run.bold = True
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        shade(hdr[i], header_fill)
    for row in rows:
        cells = t.add_row().cells
        for i, text in enumerate(row):
            cells[i].text = str(text)
    if widths:
        for r in t.rows:
            for i, w in enumerate(widths):
                r.cells[i].width = w
    doc.add_paragraph()
    return t


def build(path):
    doc = Document()
    set_base_style(doc)
    for name, size in (("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 11.5)):
        style_heading(doc, name, size, ACCENT)

    title = doc.add_heading("Kaggriculture 分工方案：人做什么，AI 做什么", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub = p(doc, "适用队伍规模 3–5 人｜配合《Kaggriculture 队员手册》使用", size=10)
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph()

    # ============================================================ 0
    h1(doc, "0  一分钟速览")

    rich(doc, [("划分依据只有一条：", True, None),
               ("AI 擅长「已有明确判据的批量劳动」，人负责「判断、线上操作和对外交付」。", False, None)])
    doc.add_paragraph()

    table(doc,
          ["交给 AI 智能体", "必须由人来做"],
          [
              ["写脚本、解析回放、跑批量实验、整理数据",
               "决定目标、判断哪条路值得做、取舍方案"],
              ["写 Agent 代码、重构、改 bug",
               "登录 Kaggle、提交、下载、接受规则"],
              ["本地对战评测、参数扫描、统计表格",
               "在浏览器里手动玩一局、建立直觉"],
              ["生成报告初稿、文档、图表",
               "最终结论签字、上台汇报、答辩"],
              ["重复性的格式转换与校对",
               "组队、分工、时间管理、队友沟通"],
          ],
          widths=[Pt(240), Pt(240)])

    rich(doc, [("一句话：", True, None),
               ("让 AI 当「执行力」，人当「决策者」。"
                "不要让 AI 替你决定做什么，也不要让人去干重复劳动。", False, None)])

    doc.add_page_break()

    # ============================================================ 1
    h1(doc, "1  为什么这样分")
    p(doc, "三类工作对应的边界：", bold=True)
    table(doc,
          ["工作类型", "特点", "归属"],
          [
              ["需要判断力的", "信息不全、要权衡取舍、错了代价大",
               "人。AI 没有你的目标函数"],
              ["必须动账号/浏览器/线下的", "Kaggle 登录、提交、下载、开会",
               "人。AI 做不了，也不该代做"],
              ["有明确判据的重复劳动", "解析、统计、批量实验、写模板代码",
               "AI。这是它性价比最高的地方"],
          ],
          widths=[Pt(120), Pt(230), Pt(130)])

    warn(doc, "关键认知：AI 的能力上限由你给它的「判据」决定。"
              "如果你只说「优化一下」，它只能瞎猜；如果你说「胜率低于 55% 就换方案」，它就能干活。"
              "所以「把判据写清楚」是人不可推卸的责任。")

    # ============================================================ 2
    h1(doc, "2  人负责的事（不要交给 AI）")

    h2(doc, "2.1  判断与决策")
    bullets(doc, [
        ("定目标。", "明确这一阶段要什么（跑通？提胜率？保住名次？），以及优先级。"),
        ("做取舍。", "多个方案冲突时决定采哪个；AI 会给你数据，但不会替你负责。"),
        ("判断改动是否值得做。", "AI 能算出「翻了 46 局败局、新增 8 局败局」，"
         "但值不值得提交，是人按风险偏好定的。"),
        ("识别方向性错误。", "例如把「多赚钱」当成目标，AI 不会主动纠正你。"),
    ])

    h2(doc, "2.2  线上操作（AI 无法代做）")
    bullets(doc, [
        ("登录 Kaggle 并接受竞赛规则。", "报名相关的操作必须人工完成。"),
        ("提交 Agent。", "每天最多 5 次，但只有最新的 2 次被追踪；"
         "最后两次提交的时机由人掌握。"),
        ("下载回放与自己的提交记录。", "涉及账号凭据，由人执行。"),
        ("确认各组队相关截止日。", "如组队合并截止日，需在竞赛页面人工核对。"),
    ])
    warn(doc, "提交是最不可逆的动作。AI 可以告诉你「这版准备好了」，"
              "但按不按提交键必须是人决定的，而且人要清楚当前在册的是哪两个版本。")

    h2(doc, "2.3  线下直觉建立")
    bullets(doc, [
        ("在浏览器里手动玩一局。", "这是建立游戏直觉最快的方式，"
         "比读任何文档都有效。只有玩过，你才能判断 AI 给的方案是否合理。"),
        ("看回放。", "重点关注：分差大小、随机种子、是否反复输给同一个对手。"),
        ("亲手跑一遍关键脚本。", "至少看一次原始输出，不要只看 AI 的转述。"
         "这一步能发现 AI 的误读。"),
    ])

    h2(doc, "2.4  对外交付")
    bullets(doc, [
        ("每周报告的最终结论。", "AI 可以写初稿，但结论要人确认并签字。"),
        ("抽签上台汇报、答辩。", "必须由人来讲，讲不清说明自己没真正理解。"),
        ("组队、内部分工、进度管理。", "纯人际事务。"),
        ("期末书面报告定稿。", "文责自负。"),
    ])

    h2(doc, "2.5  把需求说清楚（最容易被忽略）")
    p(doc, "AI 的输出质量几乎完全取决于你的输入。交办任务前，先回答这五个问题：")
    table(doc,
          ["要交代的", "反例（模糊）", "正例（可执行）"],
          [
              ["目标", "「优化一下 Agent」", "「把对本地强陪练的胜率从 46% 提到 55% 以上」"],
              ["范围", "「改改市场模块」", "「只改市场卖单的排序逻辑，其他模块不动」"],
              ["判定标准", "「看看效果好不好」", "「跑 20 个种子，胜率提升且回归集不倒退」"],
              ["输出格式", "「给我结果」", "「先给 200 字摘要，我确认后再给完整数据」"],
              ["边界", "（没说）", "「不要下载任何外部 Agent，不要动我的账号」"],
          ],
          widths=[Pt(80), Pt(180), Pt(220)])

    doc.add_page_break()

    # ============================================================ 3
    h1(doc, "3  交给 AI 智能体做的事")

    h2(doc, "3.1  代码与工程")
    table(doc,
          ["任务", "说明", "产物"],
          [
              ["实现 Agent 主逻辑", "按人给定的规则写代码，不是自己发明策略",
               "可运行、可提交的 Agent 文件"],
              ["重构与修 bug", "把逻辑拆成模块、加异常兜底、修运行时错误",
               "稳定不报错的代码"],
              ["搭本地对战环境", "让不同版本的 Agent 互相对打", "对战脚本"],
              ["写参数扫描工具", "批量试不同参数组合并汇总", "扫描脚本 + 结果表"],
              ["打包提交文件", "按竞赛要求组织压缩包结构", "提交包"],
          ],
          widths=[Pt(110), Pt(220), Pt(150)])

    h2(doc, "3.2  数据分析")
    table(doc,
          ["任务", "说明"],
          [
              ["解析回放", "把对局记录转成可统计的结构化数据"],
              ["统计高胜率动作序列", "从大量对局里找出反复出现的有效操作"],
              ["算评测指标", "胜率、胜率矩阵、平均分差、方差、最差情况、成对 delta"],
              ["做反事实验证", "固定对手轨迹、只替换我方方案、跑多个种子并对比"],
              ["剔除无效样本", "过滤掉「分高但提交久远且正在连败」的对手数据"],
          ],
          widths=[Pt(130), Pt(350)])

    h2(doc, "3.3  文档与报告")
    table(doc,
          ["任务", "说明"],
          [
              ["每周报告初稿", "按固定框架填充，人再改结论"],
              ["生成表格与图表", "把统计结果整理成可直接放进报告的表格"],
              ["格式转换", "Markdown 转 Word、整理排版"],
              ["校对与一致性检查", "检查文档里的说法是否前后矛盾"],
          ],
          widths=[Pt(130), Pt(350)])

    h2(doc, "3.4  AI 明确不要做的事")
    bullets(doc, [
        ("不要替人登录账号、提交、下载。", "这些动作由人执行。"),
        ("不要自行引入外部方案。", "在未获明确许可前，不要下载或抄袭他人的 Agent。"),
        ("不要代替人下最终结论。", "它可以给建议和依据，但结论由人负责。"),
        ("不要在没跑够样本时给结论。", "至少要多个随机种子，单局结果没有意义。"),
    ])

    doc.add_page_break()

    # ============================================================ 4
    h1(doc, "4  怎么用 AI 更省成本")

    p(doc, "同样的活，交给 AI 的方式不同，成本可能差很多。以下四条是很实际的省钱做法。")

    h2(doc, "4.1  让 AI 写脚本，人手动执行")
    bullets(doc, [
        "很多任务是批量劳动，不需要 AI 实时盯着。",
        "正确做法：让 AI 产出脚本 → 人自己跑 → 把结果给 AI 看。",
        "反例：让 AI 全程监控一次几十分钟的批量任务，token 全花在等待和日志上。",
    ])

    h2(doc, "4.2  脚本输出两份报告，先看摘要")
    bullets(doc, [
        ("精简摘要版：", "只给结论数字，几百字，先给 AI 看。"),
        ("完整详细版：", "全量数据，先存着，摘要不够用时再给。",
         ),
        "完整报告的体量通常是摘要的几十倍，而多数情况下摘要就够了。",
    ])

    h2(doc, "4.3  克制日志输出")
    bullets(doc, [
        "运行过程中的逐行日志会大量消耗 token，且几乎没有信息量。",
        "把详细日志写进文件，只在终端打印关键节点。",
    ])

    h2(doc, "4.4  用文件当交接媒介，而不是聊天记录")
    bullets(doc, [
        "把「当前方案说明」「评测结果」「待办」写成文件放在工作区。",
        "每次开新对话，让 AI 先读这几个文件，比重新口述一遍便宜且准确。",
        "好处：上下文可复用、可交接给队友、不依赖某一次对话记忆。",
    ])

    # ============================================================ 5
    h1(doc, "5  按角色分工建议（3–5 人）")

    p(doc, "下面是一个可以直接用的分工。人数少时可以一人兼多个角色，"
           "但「决策」和「执行」不要集中在同一个人身上，否则没有互相检查。")

    table(doc,
          ["角色", "人负责", "交给 AI 的"],
          [
              ["队长 / 决策",
               "定目标、排优先级、拍板采不采纳、掌握提交时机",
               "整理决策所需的对比数据"],
              ["算法 / 策略",
               "设计规则与策略、判断方向、看回放建立直觉",
               "写 Agent 代码、跑参数扫描、做反事实验证"],
              ["数据 / 评测",
               "定义评测指标与验收标准、抽样检查数据质量",
               "解析回放、算胜率矩阵与方差、批量对战"],
              ["文档 / 汇报",
               "定结论、上台讲、答辩",
               "写报告初稿、生成表格、格式转换与校对"],
          ],
          widths=[Pt(90), Pt(200), Pt(190)])

    warn(doc, "如果只有 2–3 人：优先保证「决策」和「评测」两个角色有人，"
              "文档工作可以大量交给 AI，但结论必须人签字。")

    # ============================================================ 6
    h1(doc, "6  交接模板")

    p(doc, "每次向 AI 交办任务，照这个模板填，能显著减少来回追问。建议直接复制使用。")

    for label, content in [
        ("目标", "这一轮要达成什么，可量化最好。"),
        ("范围", "只允许改动哪些文件 / 模块；明确哪些不要动。"),
        ("输入", "需要 AI 先读的文件路径，以及参考数据的位置。"),
        ("判定标准", "用什么指标判断成功，需要跑多少样本。"),
        ("输出格式", "先摘要还是直接完整数据；要表格还是文字。"),
        ("边界与禁止项", "不许动账号、不许下载外部方案、不许改配置等。"),
    ]:
        rich(doc, [(f"{label}：", True, None), (content, False, None)])

    doc.add_paragraph()
    p(doc, "示例（可直接改写）：", bold=True)
    par = doc.add_paragraph()
    par.paragraph_format.left_indent = Pt(18)
    sample = (
        "目标：把对本地强陪练的胜率从 46% 提到 55% 以上。\n"
        "范围：只改市场卖单的排序逻辑，不动农民与雇工的行动逻辑。\n"
        "输入：先读 队员手册.docx 的第 3 节和第 5 节；评测脚本在 benchmark 类文件里。\n"
        "判定标准：跑 20 个随机种子；胜率提升，且原本赢的对局不出现倒退。\n"
        "输出格式：先给 200 字摘要 + 一张胜率表，我确认后再给完整数据。\n"
        "边界：不要下载任何外部 Agent；不要动我的 Kaggle 账号；不要提交。"
    )
    run = par.add_run(sample)
    run.font.size = Pt(9.5)
    run.font.name = "Consolas"

    doc.add_paragraph()
    rich(doc, [("回报模板（AI 给结果时也照这个来）：", True, None)])
    bullets(doc, [
        "做了什么改动（一句话）",
        "跑了多少样本、对手是谁",
        "结果数字（胜率 / 分差 / 最差情况）",
        "是否达标；没达标的可能原因",
        "建议下一步（但由人决定）",
    ])

    doc.add_paragraph()
    warn(doc, "最后提醒：AI 给出的任何结论都要人能复述一遍才算数。"
              "如果讲不清「为什么这样改有效」，就不要提交，也不要在报告里写。")

    doc.save(path)
    print("wrote", path)


if __name__ == "__main__":
    build(r"D:\hiwsk\Documents\deepseek工作区\kaggriculture\人员与AI分工方案.docx")
